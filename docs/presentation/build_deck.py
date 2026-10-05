#!/usr/bin/env python3
"""Builds docs/Kids-Explorer-AI.pptx from screenshots.

Usage:  pip install python-pptx pillow
        python docs/presentation/build_deck.py [screenshots_dir]

The screenshots dir needs the web/admin PNGs (see docs/presentation/README.md);
Android screens are taken from mobile/test/store_screenshots/screenshots/.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SHOTS = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'docs', 'presentation', 'screens')
PHONE = os.path.join(ROOT, 'mobile', 'test', 'store_screenshots', 'screenshots')
OUT = os.path.join(ROOT, 'docs', 'Kids-Explorer-AI.pptx')

INK = RGBColor(0x1B, 0x15, 0x30)
INK2 = RGBColor(0x2A, 0x21, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xC3, 0xE6)
PINK = RGBColor(0xFF, 0x4F, 0xA3)
VIOLET = RGBColor(0x7C, 0x5C, 0xFF)
BLUE = RGBColor(0x1E, 0x90, 0xFF)
GREEN = RGBColor(0x22, 0xB5, 0x73)
ORANGE = RGBColor(0xFF, 0x8A, 0x3D)
YELLOW = RGBColor(0xFF, 0xC9, 0x3C)
TEAL = RGBColor(0x00, 0xC2, 0xA8)
RED = RGBColor(0xFF, 0x5D, 0x73)
TITLE_FONT = 'Arial Black'
BODY_FONT = 'Calibri'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def shot(name, phone=False):
    for ext in ('.png', '.jpg'):
        p = os.path.join(PHONE if phone else SHOTS, name + ext)
        if os.path.exists(p):
            return p
    print('missing screenshot:', name)
    return None


def bg(slide, c1=INK, c2=INK2, angle=135):
    f = slide.background.fill
    f.gradient()
    f.gradient_angle = angle
    f.gradient_stops[0].color.rgb = c1
    f.gradient_stops[1].color.rgb = c2


def blob(slide, x, y, d, color, transparency=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    _alpha(s, transparency)
    s.line.fill.background()
    return s


def _alpha(shape, transparency):
    # python-pptx has no API for fill alpha: set it on the XML
    from pptx.oxml.ns import qn
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    if sf is None:
        return
    clr = sf[0]
    for a in clr.findall(qn('a:alpha')):
        clr.remove(a)
    el = clr.makeelement(qn('a:alpha'), {'val': str(int((1 - transparency) * 100000))})
    clr.append(el)


def text(slide, x, y, w, h, s, size=18, color=WHITE, bold=False, font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def rich(slide, x, y, w, h, items, size=17, color=WHITE, gap=6):
    """items: list of (emoji_or_bullet, bold_lead, rest)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (icon, lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        for t, b, c in ((f'{icon}  ', False, color), (lead, True, YELLOW), (rest, False, color)):
            if not t:
                continue
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.name = BODY_FONT
            r.font.color.rgb = c
    return tb


def label_width(label, size=13):
    wide = sum(1 for ch in label if ord(ch) > 0x2000)
    caps = sum(1 for ch in label if ch.isupper())
    return Inches((0.3 + 0.115 * len(label) + 0.12 * wide + 0.035 * caps) * size / 13)


def pill(slide, x, y, label, color=PINK, w=None, size=13, txt=WHITE):
    w = w or label_width(label, size)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.38))
    s.adjustments[0] = 0.5
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = txt
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return w


def pills(slide, x, y, labels, colors=(PINK, VIOLET, BLUE, GREEN, ORANGE, TEAL), size=13, maxw=None):
    cx, cy = x, y
    for i, l in enumerate(labels):
        w = label_width(l, size)
        if maxw and cx + w > x + maxw:
            cx, cy = x, cy + Inches(0.48)
        pill(slide, cx, cy, l, colors[i % len(colors)], w, size)
        cx += w + Inches(0.12)
    return cy


def card(slide, x, y, w, h, color=WHITE, transparency=0.9, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.12
    s.fill.solid()
    s.fill.fore_color.rgb = color
    _alpha(s, transparency)
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def image(slide, path, x, y, w=None, h=None, frame=True, radius=0.04):
    if not path:
        return None
    if frame:
        pad = Inches(0.07)
        fw = w if w else None
        pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
        fr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x - pad, y - pad, pic.width + 2 * pad, pic.height + 2 * pad)
        fr.adjustments[0] = radius
        fr.fill.solid()
        fr.fill.fore_color.rgb = WHITE
        _alpha(fr, 0.82)
        fr.line.fill.background()
        # put frame behind picture
        slide.shapes._spTree.remove(fr._element)
        slide.shapes._spTree.insert(2, fr._element)
        del fw
        return pic
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def phone(slide, path, x, y, h):
    """Draws a phone bezel around an Android screenshot."""
    if not path:
        return
    from PIL import Image

    iw, ih = Image.open(path).size
    sh = h
    sw = int(h * iw / ih)
    b = Inches(0.11)
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x - b, y - b * 2, sw + 2 * b, sh + 4 * b)
    body.adjustments[0] = 0.09
    body.fill.solid()
    body.fill.fore_color.rgb = RGBColor(0x0E, 0x0B, 0x1C)
    body.line.color.rgb = RGBColor(0x4A, 0x40, 0x7A)
    body.line.width = Pt(2)
    slide.shapes.add_picture(path, x, y, width=sw, height=sh)
    notch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + sw // 2 - Inches(0.25), y - b * 1.5, Inches(0.5), Inches(0.07))
    notch.adjustments[0] = 0.5
    notch.fill.solid()
    notch.fill.fore_color.rgb = RGBColor(0x33, 0x2B, 0x55)
    notch.line.fill.background()
    return sw


def header(slide, kicker, title, color=PINK, sub=None):
    pill(slide, Inches(0.6), Inches(0.45), kicker, color, size=12)
    text(slide, Inches(0.6), Inches(0.92), Inches(12.2), Inches(0.9), title, size=30, bold=True, font=TITLE_FONT)
    if sub:
        text(slide, Inches(0.6), Inches(1.72), Inches(11.5), Inches(0.6), sub, size=17, color=SOFT)


def footer(slide, n):
    text(slide, Inches(0.6), Inches(7.05), Inches(6), Inches(0.3), 'Kids Explorer AI  ·  Learn • Play • Imagine', size=10, color=SOFT)
    text(slide, Inches(12.2), Inches(7.05), Inches(0.6), Inches(0.3), str(n), size=10, color=SOFT, align=PP_ALIGN.RIGHT)


def new(c1=INK, c2=INK2, blobs=True):
    s = prs.slides.add_slide(BLANK)
    bg(s, c1, c2)
    if blobs:
        blob(s, Inches(-1.5), Inches(-1.8), Inches(4.5), PINK, 0.82)
        blob(s, Inches(10.8), Inches(4.6), Inches(4.2), VIOLET, 0.8)
        blob(s, Inches(6.5), Inches(6.4), Inches(2.4), YELLOW, 0.88)
    return s


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


n = 0

# 1 ─ Title ────────────────────────────────────────────────────────────────
s = new(RGBColor(0x2B, 0x0B, 0x4F), RGBColor(0x0B, 0x0B, 0x2B))
blob(s, Inches(7.4), Inches(-2.2), Inches(7), PINK, 0.78)
blob(s, Inches(9.8), Inches(3.4), Inches(5), BLUE, 0.8)
text(s, Inches(0.7), Inches(0.8), Inches(7), Inches(0.5), '🚀  A free AI learning world for kids aged 2–10+', size=18, color=YELLOW, bold=True)
text(s, Inches(0.7), Inches(1.45), Inches(6.2), Inches(2.2), ['Kids Explorer', 'AI'], size=54, bold=True, font=TITLE_FONT)
text(s, Inches(0.7), Inches(3.55), Inches(6.2), Inches(1.2), 'Alphabets • Stories • Brain games • Talking AI buddy • Cartoons — in English, বাংলা and more. No sign-up. Safe. Colorful.', size=19, color=SOFT)
pills(s, Inches(0.7), Inches(4.85), ['🌐 Website', '📱 Android app', '🛠️ Admin panel'], (PINK, GREEN, VIOLET), size=15)
pills(s, Inches(0.7), Inches(5.45), ['0 sign-ups', '5 languages', '15 explore topics', '5 brain games', '100% free AI buddy'], (INK2,) * 5, size=12)
image(s, shot('web-home'), Inches(7.3), Inches(1.15), w=Inches(4.9))
phone(s, shot('03-home-space', True), Inches(10.9), Inches(3.35), Inches(3.55))
n += 1
footer(s, n)
notes(s, 'Kids Explorer AI is one ecosystem: a website, a Google Play Android app and an admin dashboard that controls both.')

# 2 ─ Vision ───────────────────────────────────────────────────────────────
s = new()
header(s, 'WHY', 'Learning that feels like a game — safe & free', PINK)
cols = [
    ('🎮', 'Zero barrier', 'No account, no email, no payment. Type a name and play.', PINK),
    ('🌏', 'Multilingual', 'English, বাংলা, Español, Français, हिन्दी + any language the admin adds.', BLUE),
    ('🧠', 'Learn by doing', 'See → hear → trace → say. Every activity is a 2–5 minute "level".', GREEN),
    ('🛡️', 'Safe by design', 'Data stays on the device. Child-directed ads. Parental gates.', ORANGE),
]
for i, (e, t, d, c) in enumerate(cols):
    x = Inches(0.6 + i * 3.1)
    card(s, x, Inches(2.5), Inches(2.85), Inches(3.6), c, 0.15)
    text(s, x + Inches(0.3), Inches(2.75), Inches(2.3), Inches(1), e, size=48)
    text(s, x + Inches(0.3), Inches(3.85), Inches(2.4), Inches(0.6), t, size=22, bold=True, font=TITLE_FONT)
    text(s, x + Inches(0.3), Inches(4.55), Inches(2.35), Inches(1.5), d, size=15)
text(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.5), 'Design informed by leading early-learning apps (Khan Academy Kids, Lingokids, Duolingo ABC): age-adaptive, short lessons, levels & rewards.', size=14, color=SOFT)
n += 1
footer(s, n)

# 3 ─ Ecosystem ────────────────────────────────────────────────────────────
s = new()
header(s, 'ECOSYSTEM', 'Three apps · one content brain', VIOLET)
blocks = [
    ('🌐', 'Kids Website', 'React 18 · Vite · Framer Motion · Web Speech API', 'feature/web-app', PINK),
    ('📱', 'Android App', 'Flutter · AdMob (Families) · on-device TTS & speech', 'feature/mobile-app', GREEN),
    ('🛠️', 'Admin Panel', 'React · Tailwind CSS · Recharts · Supabase', 'feature/admin-panel', VIOLET),
]
for i, (e, t, d, br, c) in enumerate(blocks):
    x = Inches(0.6 + i * 4.15)
    card(s, x, Inches(2.2), Inches(3.9), Inches(2.3), c, 0.12)
    text(s, x + Inches(0.3), Inches(2.35), Inches(1), Inches(0.9), e, size=40)
    text(s, x + Inches(1.3), Inches(2.5), Inches(2.5), Inches(0.6), t, size=22, bold=True, font=TITLE_FONT)
    text(s, x + Inches(0.3), Inches(3.35), Inches(3.4), Inches(0.8), d, size=15)
    pill(s, x + Inches(0.3), Inches(3.95), '🌿 ' + br, INK, size=12)
card(s, Inches(0.6), Inches(4.85), Inches(12.1), Inches(1.95), WHITE, 0.92)
rich(s, Inches(0.9), Inches(5.0), Inches(11.6), Inches(1.8), [
    ('🗄️', 'Supabase backend (optional, free tier): ', 'Postgres JSONB content + Row-Level Security, media storage, anonymous insert-only analytics, optional AI edge function.'),
    ('📦', 'Offline-first: ', 'every app ships the bundled content, caches the latest admin version and refreshes in the background.'),
    ('🔁', 'No backend? ', 'Admin demo mode exports one content.json that the website and app can load from any static host.'),
], size=15)
n += 1
footer(s, n)

# 4 ─ Onboarding ───────────────────────────────────────────────────────────
s = new()
header(s, 'KID JOURNEY', 'Onboarding is the first game: 4 levels', PINK, 'An animated robot asks (and speaks) each question — name → age → hero → magic colors.')
steps = [('1', '✍️ Name', 'Live "Hello Mahi!"'), ('2', '🎈 Age', '2–4 · 5–7 · 8–10+'), ('3', '🦸 Hero', 'Boy · Girl · Superhero · Robot · Cosmic'), ('4', '🎨 Theme', '"★ for you" picked by hero')]
for i, (num, t, d) in enumerate(steps):
    y = Inches(2.55 + i * 1.02)
    c = [PINK, ORANGE, GREEN, VIOLET][i]
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), y, Inches(0.7), Inches(0.7))
    o.fill.solid()
    o.fill.fore_color.rgb = c
    o.line.fill.background()
    o.text_frame.text = num
    o.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    o.text_frame.paragraphs[0].runs[0].font.bold = True
    o.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, Inches(1.5), y - Inches(0.02), Inches(3.5), Inches(0.4), t, size=20, bold=True)
    text(s, Inches(1.5), y + Inches(0.36), Inches(3.5), Inches(0.4), d, size=14, color=SOFT)
image(s, shot('web-onboarding-theme'), Inches(5.0), Inches(2.45), w=Inches(5.4))
phone(s, shot('01-onboarding', True), Inches(10.95), Inches(2.45), Inches(4.2))
text(s, Inches(0.6), Inches(6.55), Inches(4.4), Inches(0.4), 'Saved only on the device (LocalStorage / SharedPreferences).', size=13, color=SOFT)
n += 1
footer(s, n)

# 5 ─ Adventure map ────────────────────────────────────────────────────────
s = new()
header(s, 'MOTIVATION', 'A map that keeps kids coming back', ORANGE)
image(s, shot('web-home'), Inches(0.7), Inches(2.0), w=Inches(7.2))
rich(s, Inches(8.4), Inches(2.0), Inches(4.4), Inches(4.6), [
    ('🗺️', 'Levels on a map ', '— admin chooses which modules appear and in which order.'),
    ('⭐', 'Stars → levels ', '(every 10 stars) with an animated progress bar.'),
    ('🎯', 'Daily missions ', 'rotate every day across modules — spaced practice.'),
    ('🔥', 'Streaks, 🎁 stickers ', '(every 5 ⭐) and 🏅 7 badges.'),
    ('🎉', 'Confetti, sounds & haptics ', 'on every win.'),
    ('⏱️', 'Healthy limits: ', 'daily screen time + wiggle-break reminders.'),
], size=16)
n += 1
footer(s, n)

# 6 ─ Alphabet ─────────────────────────────────────────────────────────────
s = new(RGBColor(0x3A, 0x0B, 0x3F), INK)
header(s, 'ALPHABET & PHONICS', 'See it · hear it · trace it · say it', PINK)
image(s, shot('web-alphabet-letter'), Inches(0.7), Inches(2.0), w=Inches(6.4))
phone(s, shot('04-alphabet-bangla', True), Inches(7.55), Inches(2.0), Inches(4.5))
rich(s, Inches(10.15), Inches(2.0), Inches(2.9), Inches(4.6), [
    ('🇬🇧', 'English ', '26'),
    ('🇧🇩', 'বাংলা ', '11 + 39 + ০-৯'),
    ('🇪🇸', 'Español ', '27'),
    ('🇫🇷', 'Français ', '26'),
    ('🇮🇳', 'हिन्दी ', '12 + 31'),
    ('➕', 'Any language ', 'added by admin'),
    ('✏️', 'Tracing ', 'with sparkle particles'),
    ('🎤', 'Pronunciation ', '0-3 ⭐ + slow syllables'),
], size=15, gap=4)
n += 1
footer(s, n)

# 7 ─ Stories ──────────────────────────────────────────────────────────────
s = new(RGBColor(0x0B, 0x2B, 0x4F), INK)
header(s, 'STORYTELLER', 'Flip-book stories that read themselves', BLUE, 'Narration with word-by-word highlighting → moral of the story → quiz → stars.')
image(s, shot('web-story'), Inches(0.7), Inches(2.5), w=Inches(6.6))
phone(s, shot('05-story', True), Inches(7.8), Inches(2.4), Inches(4.3))
rich(s, Inches(10.2), Inches(2.5), Inches(2.9), Inches(4.2), [
    ('📖', '3D page flip ', '& swipe'),
    ('🔊', 'Read to me ', '+ auto-play'),
    ('🖍️', 'Highlight ', 'each spoken word'),
    ('💡', 'Moral card ', 'after every story'),
    ('❓', 'Quiz ', 'to check understanding'),
    ('🛠️', 'Admin builder: ', 'pages, art, audio'),
], size=15)
n += 1
footer(s, n)

# 8 ─ Games ────────────────────────────────────────────────────────────────
s = new(RGBColor(0x0B, 0x3B, 0x2A), INK)
header(s, 'BRAIN GAMES', 'Five games that grow with the child', GREEN)
image(s, shot('web-sudoku'), Inches(0.7), Inches(2.0), w=Inches(6.3))
phone(s, shot('07-games', True), Inches(7.5), Inches(2.0), Inches(4.5))
rich(s, Inches(10.1), Inches(2.0), Inches(3.0), Inches(4.8), [
    ('🔢', 'Kids Sudoku ', '3×3 pictures → 4×4 → 6×6, hints, conflict glow'),
    ('🃏', 'Memory Match ', '3 → 8 pairs by age'),
    ('🔺', 'What comes next? ', 'AB/ABC patterns & number sequences, drag & drop'),
    ('➕', 'Number Pop ', 'count → add → × tables'),
    ('🕵️', 'Odd One Out ', 'cross-topic reasoning'),
], size=14)
n += 1
footer(s, n)

# 9 ─ Explore ──────────────────────────────────────────────────────────────
s = new(RGBColor(0x4F, 0x2B, 0x0B), INK)
header(s, 'EXPLORE THE WORLD', '15 topics · 141 things to discover', ORANGE)
image(s, shot('web-explore'), Inches(0.7), Inches(2.05), w=Inches(6.3))
phone(s, shot('06-explore-vehicles', True), Inches(7.5), Inches(2.0), Inches(4.5))
pills(s, Inches(10.05), Inches(2.05), ['🦁 Animals', '🦜 Birds', '🍎 Fruits', '🚗 Vehicles', '🚢 Ships', '✈️ Planes', '🔷 Shapes', '🌈 Colors', '🔢 Numbers', '🌍 Countries', '🪐 Space', '🧍 Body', '🧑‍⚕️ Jobs', '😊 Feelings', '🪥 Habits'], size=12, maxw=Inches(3.0))
text(s, Inches(10.05), Inches(6.05), Inches(3.0), Inches(1.0), 'Cars drive, ships sail, planes fly. English + বাংলা names, facts, "Find it!" game.', size=12, color=SOFT)
n += 1
footer(s, n)

# 10 ─ Talking Buddy ───────────────────────────────────────────────────────
s = new(RGBColor(0x1B, 0x14, 0x46), RGBColor(0x06, 0x0B, 0x1F))
header(s, '★ STAR FEATURE', 'Talking Buddy: an AI friend for speaking', YELLOW, None)
pill(s, Inches(2.75), Inches(0.45), '100% FREE · ON-DEVICE', GREEN, size=12)
image(s, shot('web-buddy'), Inches(0.7), Inches(1.95), w=Inches(5.7))
phone(s, shot('09-buddy', True), Inches(6.85), Inches(1.95), Inches(4.75))
steps = [
    ('🛡️', 'Safety filter first', PINK),
    ('✏️', 'Grammar & dialect: "i is" → "I am"', ORANGE),
    ('👶', 'Baby talk → formal word + meaning', GREEN),
    ('🙏', '"gimme water" → "May I please have…?"', BLUE),
    ('❓', 'Punctuation & capitals', VIOLET),
    ('💬', 'Answers in voice + text', TEAL),
]
for i, (e, t, c) in enumerate(steps):
    y = Inches(1.95 + i * 0.66)
    card(s, Inches(9.45), y, Inches(3.55), Inches(0.58), c, 0.25)
    text(s, Inches(9.58), y + Inches(0.05), Inches(3.35), Inches(0.52), f'{e}  {t}', size=11, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(9.45), Inches(5.95), Inches(3.6), Inches(0.9), '5 voices (Robot, Bunny, Bear, Astronaut, Fairy) × 5 languages. Optional cloud AI via Gemini/OpenAI edge function.', size=12, color=SOFT)
n += 1
footer(s, n)

# 11 ─ AI Lab ──────────────────────────────────────────────────────────────
s = new(RGBColor(0x0B, 0x0B, 0x2B), RGBColor(0x2A, 0x0B, 0x4F))
header(s, 'FUTURE SKILLS', 'AI Lab — kids discover how AI thinks', TEAL)
image(s, shot('web-ailab'), Inches(0.7), Inches(2.0), w=Inches(6.6))
phone(s, shot('10-ai-lab', True), Inches(7.75), Inches(2.0), Inches(4.5))
rich(s, Inches(10.35), Inches(2.0), Inches(2.75), Inches(4.8), [
    ('💡', 'What is AI? ', '5 animated steps'),
    ('✏️', 'Draw & guess ', '— live confidence bars'),
    ('🧑‍🏫', 'Teach the Robot ', '— training data & k-NN'),
    ('🔮', 'Can AI predict? ', '— try to trick it!'),
    ('🤝', 'Message: ', 'AI is a helper made by people; it can be wrong.'),
], size=14)
n += 1
footer(s, n)

# 12 ─ Themes & gender ─────────────────────────────────────────────────────
s = new()
header(s, 'PERSONALISATION', 'Colors that follow the child — or the admin', PINK, 'Theme modes: Kid chooses · Automatic by gender/hero · One theme for everyone (events).')
phone(s, shot('02-home-bubblegum', True), Inches(0.75), Inches(2.5), Inches(4.2))
phone(s, shot('03-home-space', True), Inches(3.05), Inches(2.5), Inches(4.2))
phone(s, shot('08-sudoku', True), Inches(5.35), Inches(2.5), Inches(4.2))
image(s, shot('admin-themes'), Inches(7.75), Inches(2.55), w=Inches(5.1))
text(s, Inches(7.75), Inches(5.85), Inches(5.1), Inches(1), '6 editable themes · gender → theme mapping · heroes · fonts · live preview — applied remotely to web & Android.', size=13, color=SOFT)
n += 1
footer(s, n)

# 13 ─ Cartoons & safety ───────────────────────────────────────────────────
s = new()
header(s, 'SAFE CARTOONS', 'Age-curated video hub', TEAL)
image(s, shot('admin-videos'), Inches(0.7), Inches(2.0), w=Inches(6.8))
rich(s, Inches(8.0), Inches(2.0), Inches(4.9), Inches(4.6), [
    ('🎯', 'Filtered by age group ', 'chosen in onboarding'),
    ('🔗', 'Paste any YouTube link ', '— ID & thumbnail detected, preview plays'),
    ('🕶️', 'Privacy-enhanced embeds ', '(youtube-nocookie), related videos limited'),
    ('🚫', 'No outbound links: ', 'tap shield over the title bar'),
    ('🗂️', 'Categories: ', 'Songs, Alphabet, Stories, Science, Math…'),
    ('🎬', 'Also Vimeo & MP4 ', 'on the website'),
], size=16)
n += 1
footer(s, n)

# 14 ─ Admin dashboard ─────────────────────────────────────────────────────
s = new(RGBColor(0x14, 0x10, 0x28), INK2)
header(s, 'ADMIN PANEL', 'One dashboard for web + Android', VIOLET)
image(s, shot('admin-dashboard'), Inches(0.7), Inches(2.0), w=Inches(7.6))
rich(s, Inches(8.8), Inches(2.0), Inches(4.2), Inches(4.7), [
    ('📊', 'Daily sessions ', 'Web vs Android'),
    ('📖', 'Top stories, ', 'popular games, most-watched videos'),
    ('🎂', 'Age, theme, platform ', 'and language split'),
    ('📈', 'Reports ', 'with filters & CSV export'),
    ('🔒', 'Anonymous only: ', 'no names, no device IDs'),
    ('👥', 'Roles: ', 'owner · editor · viewer'),
], size=16)
n += 1
footer(s, n)

# 15 ─ Content engine ──────────────────────────────────────────────────────
s = new(RGBColor(0x14, 0x10, 0x28), INK2)
header(s, 'CONTENT ENGINE', 'Change everything without an app update', VIOLET)
image(s, shot('admin-story-builder'), Inches(0.7), Inches(2.0), w=Inches(6.0))
image(s, shot('admin-buddy'), Inches(7.1), Inches(2.0), w=Inches(5.6))
pills(s, Inches(0.7), Inches(6.0), ['📖 Storybook Builder', '📺 Video Curator', '🔤 Alphabet & Words', '🦁 Explore Topics', '❓ Quiz', '🤖 Buddy rules + test console', '⚙️ Features & level order', '💾 Import / Export'], size=12, maxw=Inches(12))
n += 1
footer(s, n)

# 16 ─ AdMob & Play ────────────────────────────────────────────────────────
s = new(RGBColor(0x0B, 0x2A, 0x1F), INK)
header(s, 'MONETISATION, RESPONSIBLY', 'AdMob ready for Google Play Families', GREEN)
image(s, shot('admin-admob'), Inches(0.7), Inches(2.0), w=Inches(6.6))
checks = [
    'Child-directed requests, max rating G',
    'Non-personalised ads, no remarketing',
    'AD_ID permission removed',
    'Banners on menus only, capped interstitials',
    'Rewarded ads behind the parental gate',
    'Remote kill-switch from the admin',
    'No accounts, data stays on device',
]
for i, c in enumerate(checks):
    y = Inches(2.0 + i * 0.62)
    card(s, Inches(7.75), y, Inches(5.25), Inches(0.52), GREEN, 0.3)
    text(s, Inches(7.9), y + Inches(0.06), Inches(5.0), Inches(0.45), '✅  ' + c, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
n += 1
footer(s, n)

# 17 ─ Architecture ────────────────────────────────────────────────────────
s = new()
header(s, 'UNDER THE HOOD', 'Architecture & tech stack', BLUE)
nodes = [
    (Inches(0.7), Inches(2.3), '🛠️ Admin Panel', 'React · Tailwind · Supabase Auth', VIOLET),
    (Inches(4.85), Inches(2.3), '🗄️ Supabase', 'Postgres JSONB + RLS · Storage · Events', BLUE),
    (Inches(9.0), Inches(1.75), '🌐 Website', 'React · Vite · Framer Motion', PINK),
    (Inches(9.0), Inches(3.2), '📱 Android', 'Flutter · AdMob · TTS/STT', GREEN),
    (Inches(4.85), Inches(4.45), '🧠 Buddy AI (optional)', 'Gemini / OpenAI-compatible', ORANGE),
    (Inches(0.7), Inches(4.45), '📦 shared/content', 'JSON seed → every app', TEAL),
]
for x, y, t, d, c in nodes:
    card(s, x, y, Inches(3.6), Inches(1.15), c, 0.15)
    text(s, x + Inches(0.2), y + Inches(0.12), Inches(3.3), Inches(0.4), t, size=17, bold=True)
    text(s, x + Inches(0.2), y + Inches(0.58), Inches(3.3), Inches(0.5), d, size=13, color=SOFT)
for (x1, y1, x2, y2) in [(4.3, 2.85, 4.85, 2.85), (8.45, 2.7, 9.0, 2.3), (8.45, 3.0, 9.0, 3.75), (6.65, 3.45, 6.65, 4.45), (2.5, 4.45, 2.5, 3.45)]:
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = YELLOW
    ln.line.width = Pt(3)
rich(s, Inches(0.7), Inches(5.85), Inches(12.2), Inches(1.1), [
    ('⚡', 'Offline-first: ', 'seed → cache → refresh.  '),
    ('🔐', 'RLS: ', 'apps read published content & insert anonymous events only; admins write.'),
], size=14)
n += 1
footer(s, n)

# 18 ─ Quality ─────────────────────────────────────────────────────────────
s = new()
header(s, 'QUALITY', 'Tested, linted and documented', TEAL)
stats = [('52', 'Flutter tests', 'engine parity, games, every screen rendered at phone size, onboarding flow', GREEN), ('14', 'Web unit tests', 'buddy engine, sudoku, AI-lab algorithms, content', PINK), ('0', 'Lint issues', 'flutter analyze · ESLint (web + admin)', VIOLET), ('8', 'Guides', 'README per app, build guide (Win/mac), Play Store, privacy, admin', ORANGE)]
for i, (big, t, d, c) in enumerate(stats):
    x = Inches(0.6 + i * 3.1)
    card(s, x, Inches(2.3), Inches(2.85), Inches(3.4), c, 0.15)
    text(s, x + Inches(0.3), Inches(2.45), Inches(2.4), Inches(1.2), big, size=60, bold=True, font=TITLE_FONT, color=YELLOW)
    text(s, x + Inches(0.3), Inches(3.75), Inches(2.4), Inches(0.5), t, size=19, bold=True)
    text(s, x + Inches(0.3), Inches(4.3), Inches(2.4), Inches(1.4), d, size=13, color=SOFT)
text(s, Inches(0.6), Inches(6.05), Inches(12), Inches(0.6), 'Store screenshots of the real Android screens are generated automatically: flutter test --tags screenshots --run-skipped --update-goldens', size=14, color=SOFT)
n += 1
footer(s, n)

# 19 ─ Get started ─────────────────────────────────────────────────────────
s = new()
header(s, 'GET STARTED', 'From clone to Play Store', PINK, 'Works on Windows and macOS — full guide in docs/BUILD_GUIDE.md')
code = [
    ('🌐 Website', 'cd web && npm install && npm run dev', PINK),
    ('🛠️ Admin', 'cd admin && npm install && npm run dev', VIOLET),
    ('📱 Android', 'cd mobile && flutter pub get && flutter run', GREEN),
    ('📦 Play Store', 'flutter build appbundle --release -PadmobAppId=…', ORANGE),
]
for i, (t, c, col) in enumerate(code):
    y = Inches(2.5 + i * 1.0)
    pill(s, Inches(0.7), y + Inches(0.12), t, col, w=Inches(1.95), size=12)
    card(s, Inches(2.8), y, Inches(9.9), Inches(0.66), RGBColor(0x0A, 0x08, 0x18), 0.1)
    text(s, Inches(3.05), y + Inches(0.13), Inches(9.5), Inches(0.5), c, size=17, font='Consolas', color=RGBColor(0x9C, 0xFF, 0xD6))
text(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.4), 'Next: offline voice packs · handwriting recognition for tracing · classroom mode · iOS', size=14, color=SOFT)
n += 1
footer(s, n)

# 20 ─ Thanks ──────────────────────────────────────────────────────────────
s = new(RGBColor(0x2B, 0x0B, 0x4F), RGBColor(0x0B, 0x0B, 0x2B))
blob(s, Inches(4.5), Inches(1.0), Inches(4.5), PINK, 0.85)
text(s, Inches(0.5), Inches(2.3), Inches(12.3), Inches(1.4), 'Learn • Play • Imagine', size=54, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
text(s, Inches(0.5), Inches(3.75), Inches(12.3), Inches(0.8), '🚀 🦁 📖 🧩 🤖 🌈', size=40, align=PP_ALIGN.CENTER)
text(s, Inches(0.5), Inches(4.75), Inches(12.3), Inches(0.6), 'Kids Explorer AI — website · Android app · admin panel', size=20, color=SOFT, align=PP_ALIGN.CENTER)
n += 1
footer(s, n)

prs.save(OUT)
print('saved', OUT, 'slides:', n)
